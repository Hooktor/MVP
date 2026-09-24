from decimal import Decimal, InvalidOperation
from datetime import timedelta, date, datetime
from pathlib import Path
from zipfile import BadZipFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q, Count, Avg, F
from django.http import HttpResponse, FileResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import resolve, Resolver404
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_POST
from .models import *
from .forms import ExpertForm, ExpertRequestForm, EvaluationForm
from .selectors import *
from .services import submit_request, create_solicitation, create_direct_solicitation, create_broadcast_proposal, register_purchase_order, register_delivery_report, create_evaluation, accept_solicitation, refuse_solicitation, audit

def listing(request, qs, template, partial, **ctx):
    page = Paginator(qs, 12).get_page(request.GET.get('page'))
    params=request.GET.copy(); params.pop('page',None)
    ctx.update(page=page, q=request.GET.get('q',''), query=params.urlencode(), partial=partial)
    template=partial if request.headers.get('HX-Request') == 'true' else template
    return render(request,template,ctx)

def guard(user, roles):
    if role(user) not in roles: raise PermissionDenied

def request_writable(user,obj):
    return role(user) == 'ADMIN_OMEA' or (role(user)=='ADMIN_FILIALE' and obj.subsidiary_id==subsidiary_id(user))

def with_deadline(s):
    remaining=(s.deadline-timezone.now()).total_seconds()
    s.deadline_level='expired' if remaining<=0 else 'critical' if remaining<6*3600 else 'warning' if remaining<24*3600 else 'normal'
    s.remaining='Échéance dépassée' if remaining<=0 else f'{int(remaining//3600)} h {int(remaining%3600//60):02d} min'
    return s

@login_required
def dashboard(request):
    reqs=requests_for(request.user); experts=experts_for(request.user); sols=solicitations_for(request.user); missions=missions_for(request.user)
    counts=dict(reqs.values('status').annotate(total=Count('id')).values_list('status','total'))
    response=sols.filter(status__in=['ACCEPTED','REFUSED'])
    total=response.count(); within=response.filter(responded_at__lte=F('deadline')).count()
    metrics=[('Demandes actives',reqs.exclude(status__in=['CLOSED','CANCELLED','DRAFT']).count(),'En cours de sourcing','file-description','request_list'),
      ('Pré-accords attendus',sols.filter(status='PENDING').count(),'Réponse sous 48 heures','clock','solicitation_list'),
      ('SLA respecté',f'{round(within/total*100)} %' if total else '—','Sur les réponses reçues','shield-check','reports'),
      ('Experts disponibles',experts.filter(availability='AVAILABLE',validation_status='APPROVED',is_active=True).count(),'Validés par OMEA','users','matching'),
      ('Missions en cours',missions.filter(status='IN_PROGRESS').count(),'Mobilisations actives','briefcase','mission_list')]
    pipeline=[('Matching',counts.get('WAITING_MATCHING',0),'blue'),('Pré-accord',counts.get('WAITING_PRE_AGREEMENT',0),'orange'),('PO attendu',counts.get('WAITING_PO',0),'orange'),('Mission',counts.get('IN_PROGRESS',0),'blue'),('Clôturées',counts.get('CLOSED',0),'green')]
    actions=[]
    def add_action(kind, icon, title, detail, subsidiary, due, url, pk, action, tone, progress, progress_label):
        actions.append({'kind':kind,'icon':icon,'title':title,'detail':detail,'subsidiary':subsidiary,'due':due,
          'url':url,'pk':pk,'action':action,'tone':tone,'progress':progress,'progress_label':progress_label})
    for expert_request in reqs.filter(status='DRAFT',requester=request.user).order_by('-created_at'):
        add_action('Demande','file-alt','Finaliser la demande',expert_request.reference,expert_request.subsidiary.name,'Brouillon','request_detail',expert_request.pk,'Ouvrir','blue',20,'1/5')
    for expert_request in reqs.filter(status='WAITING_MATCHING',requester=request.user).order_by('-created_at'):
        add_action('Demande','search','Suivre les propositions',expert_request.reference,expert_request.subsidiary.name,'Recherche en cours','request_detail',expert_request.pk,'Voir','blue',40,'2/5')
    for solicitation in sols.filter(status='PENDING').select_related('request','expert__subsidiary').order_by('deadline'):
        if role(request.user)=='ADMIN_FILIALE' and solicitation.expert.subsidiary_id==subsidiary_id(request.user):
            detail=f'{solicitation.request.reference} · {solicitation.expert.full_name}' if solicitation.request_id else f'{solicitation.subject} · {solicitation.expert.full_name}'
            kind='Pré-accord' if solicitation.request_id else 'Sollicitation directe'
            add_action(kind,'paper-plane','Répondre à la sollicitation',detail,solicitation.expert.subsidiary.name,with_deadline(solicitation).remaining,'solicitation_detail',solicitation.pk,'Répondre','mint',50,'2/4')
    for expert_request in reqs.filter(status='WAITING_PO',requester=request.user).order_by('created_at'):
        add_action('Bon de commande','file-invoice','Ajouter le bon de commande',expert_request.reference,expert_request.subsidiary.name,'PO attendu','request_detail',expert_request.pk,'Ajouter le PO','peach',60,'3/5')
    for mission in missions.filter(request__status='IN_PROGRESS').order_by('-start_date'):
        add_action('Procès-verbal','file-signature','Préparer le procès-verbal',mission.request.reference,mission.supplier_subsidiary.name,'Mission en cours','mission_detail',mission.pk,'Ouvrir','lavender',80,'4/5')
    for mission in missions.filter(request__status='WAITING_EVALUATION',request__requester=request.user).order_by('planned_end_date'):
        add_action('Évaluation','star','Évaluer la mission',f'{mission.request.reference} · {mission.expert.full_name}',mission.request.subsidiary.name,'PV reçu','mission_detail',mission.pk,'Évaluer','peach',90,'4/5')
    return render(request,'dashboard.html',{'title':'Vue d’ensemble','section':'dashboard','metrics':metrics,'pipeline':pipeline,
      'recent_requests':reqs.order_by('-created_at')[:5],'pending': [with_deadline(s) for s in sols.filter(status='PENDING').order_by('deadline')[:4]],
      'events':audit_for(request.user)[:5], 'available':experts.filter(availability='AVAILABLE',validation_status='APPROVED')[:3],'action_items':actions[:8]})

def expert_filters(request,qs):
    q=request.GET.get('q','').strip()
    if q: qs=qs.filter(Q(first_name__icontains=q)|Q(last_name__icontains=q)|Q(job_title__icontains=q)|Q(skills__name__icontains=q)|Q(bio__icontains=q)|Q(references__description__icontains=q))
    for param,field in [('subsidiary','subsidiary_id'),('cluster','subsidiary__cluster_id'),('domain','domains__id'),('skill','skills__id'),('certification','certifications__id'),('language','languages__id')]:
        value=request.GET.get(param,'')
        if value.isdigit(): qs=qs.filter(**{field:value})
    if request.GET.get('availability') in dict(Expert.AVAILABILITY): qs=qs.filter(availability=request.GET['availability'])
    for param,field in [('rate','daily_rate__lte'),('score','average_score__gte')]:
        try:
            v=Decimal(request.GET.get(param,''))
            if v.is_finite() and v>=0: qs=qs.filter(**{field:v})
        except InvalidOperation: pass
    return qs.distinct().order_by('last_name','first_name')

def filter_context():
    return {'subsidiaries':Subsidiary.objects.all(),'domains':ExpertiseDomain.objects.all(),'clusters':Cluster.objects.all(),'skills':Skill.objects.all(),'certifications':Certification.objects.all(),'languages':Language.objects.all()}

@login_required
def expert_list(request):
    return listing(request,expert_filters(request,experts_for(request.user)),'expert_list.html','partials/experts_table.html',title='Référentiel experts',section='expert_list',**filter_context())

@login_required
def matching(request):
    qs=expert_filters(request,experts_for(request.user,matching=True))
    eligible=requests_for(request.user).filter(status__in=['DRAFT','WAITING_MATCHING'])
    selected=get_object_or_404(eligible,pk=request.GET['request_id']) if request.GET.get('request_id','').isdigit() else None
    page=Paginator(qs,12).get_page(request.GET.get('page'))
    required=list(selected.required_skills.all()) if selected else []
    for expert in page.object_list:
        criteria=[]
        if selected:
            if selected.domain_id: criteria.append((selected.domain.name,30,selected.domain_id in {d.pk for d in expert.domains.all()}))
            for skill in required: criteria.append((skill.name,40/len(required),skill.pk in {s.pk for s in expert.skills.all()}))
            if selected.required_language_id: criteria.append((selected.required_language.name,10,selected.required_language_id in {l.pk for l in expert.languages.all()}))
            if selected.max_daily_rate is not None: criteria.append(('Budget',10,expert.daily_rate<=selected.max_daily_rate))
            if selected.desired_start: criteria.append(('Date de disponibilité',10,bool(expert.available_from and expert.available_from<=selected.desired_start)))
        expert.match_score=round(sum(w for label,w,ok in criteria if ok)/sum(w for label,w,ok in criteria)*100) if criteria else None
        expert.match_reasons=[label for label,w,ok in criteria if ok]
    params=request.GET.copy(); params.pop('page',None)
    return render(request,'partials/experts_grid.html' if request.headers.get('HX-Request')=='true' else 'matching.html',dict(page=page,selected_request=selected,eligible_requests=eligible,query=params.urlencode(),title='Recherche d’experts',section='matching',**filter_context()))

@login_required
def expert_detail(request,pk):
    qs=experts_for(request.user)
    if role(request.user)=='ADMIN_FILIALE':
        qs=Expert.objects.filter(Q(subsidiary_id=subsidiary_id(request.user))|Q(is_active=True,validation_status='APPROVED')).select_related('subsidiary').prefetch_related('skills','domains','certifications','verticals')
    obj=get_object_or_404(qs,pk=pk)
    return render(request,'expert_detail.html',{'expert':obj,'title':obj.full_name,'section':'expert_list',
      'can_edit':role(request.user)=='ADMIN_OMEA' or (role(request.user)=='ADMIN_FILIALE' and obj.subsidiary_id==subsidiary_id(request.user)),
      'can_approve':role(request.user)=='ADMIN_OMEA' and obj.validation_status=='DRAFT',
      'languages':obj.expertlanguageassignment_set.select_related('language'),'references':obj.references.all(),
      'missions':missions_for(request.user).filter(expert=obj),'evaluations':Evaluation.objects.filter(mission__in=missions_for(request.user),mission__expert=obj).select_related('evaluator'),
      'events':audit_for(request.user).filter(entity_type='Expert',entity_id=str(obj.pk))})

@login_required
@require_POST
def expert_approve(request,pk):
    guard(request.user,['ADMIN_OMEA'])
    expert=get_object_or_404(Expert,pk=pk)
    if expert.validation_status!='DRAFT':
        messages.info(request,'Cet expert a déjà été traité.')
    else:
        expert.validation_status='APPROVED'; expert.save(update_fields=['validation_status'])
        audit(request.user,'APPROVE_EXPERT',expert,'Expert approuvé et disponible pour le matching',old='DRAFT',new='APPROVED')
        messages.success(request,f'{expert.full_name} est approuvé et peut maintenant être sollicité.')
    return redirect('expert_detail',pk)

@login_required
def expert_edit(request,pk=None):
    guard(request.user,['ADMIN_FILIALE','ADMIN_OMEA'])
    obj=get_object_or_404(experts_for(request.user),pk=pk) if pk else None
    form=ExpertForm(request.POST or None,instance=obj,user=request.user)
    if request.method=='POST' and form.is_valid():
        with transaction.atomic():
            obj=form.save(); audit(request.user,'EXPERT_UPDATED',obj,'Profil expert enregistré')
        messages.success(request,'Le profil expert a été enregistré.')
        return redirect('expert_detail',obj.pk)
    return render(request,'form.html',{'form':form,'title':'Modifier le profil' if pk else 'Nouvel expert','section':'expert_list','back_url':'expert_list'})

IMPORT_HEADERS = ['Matricule *','Prénom *','Nom *','Fonction *','Code filiale','Domaines','Compétences','Certifications','Verticales','Langues (Nom:Niveau)','TJM','Devise','Disponibilité','Date disponibilité','Présentation','Statut validation']
IMPORT_TEMPLATE = Path(__file__).resolve().parent.parent / 'static' / 'downloads' / 'modele_import_experts_orbit.xlsx'

def import_values(value):
    return [part.strip() for part in str(value or '').split(',') if part and str(part).strip()]

def import_date(value):
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    parsed=parse_date(str(value or '').strip())
    if value and not parsed: raise ValueError('La date de disponibilité est invalide (format AAAA-MM-JJ attendu).')
    return parsed

def import_expert_row(values, user, line, seen_ids):
    values=[value.strip() if isinstance(value,str) else value for value in values]
    employee_id, first_name, last_name, job_title=[str(value or '').strip() for value in values[:4]]
    if not all((employee_id,first_name,last_name,job_title)):
        raise ValueError('Matricule, prénom, nom et fonction sont obligatoires.')
    if employee_id in seen_ids or Expert.objects.filter(employee_id=employee_id).exists():
        raise ValueError(f'Le matricule « {employee_id} » existe déjà.')
    requested_code=str(values[4] or '').strip().upper()
    user_subsidiary_id=subsidiary_id(user)
    if role(user)=='ADMIN_FILIALE':
        if not user_subsidiary_id: raise ValueError('Votre compte ne possède pas de filiale de rattachement.')
        subsidiary=Subsidiary.objects.get(pk=user_subsidiary_id)
        if requested_code and requested_code != subsidiary.code.upper():
            raise ValueError('Le code filiale ne correspond pas à votre périmètre.')
    else:
        if not requested_code: raise ValueError('Le code filiale est obligatoire pour un import OMEA.')
        subsidiary=Subsidiary.objects.filter(code__iexact=requested_code).first()
        if not subsidiary: raise ValueError(f'La filiale « {requested_code} » est inconnue.')
    availability=str(values[12] or 'AVAILABLE').strip().upper()
    if availability not in dict(Expert.AVAILABILITY): raise ValueError('La disponibilité est invalide.')
    validation_status=str(values[15] or 'DRAFT').strip().upper()
    if validation_status not in dict(Expert.VALIDATION): raise ValueError('Le statut de validation est invalide.')
    if role(user)!='ADMIN_OMEA': validation_status='DRAFT'
    try:
        daily_rate=Decimal(str(values[10]).replace(',','.')) if values[10] else Decimal('0')
    except InvalidOperation: raise ValueError('Le TJM doit être un nombre valide.')
    if not daily_rate.is_finite() or daily_rate<0: raise ValueError('Le TJM doit être positif ou nul.')
    currency=str(values[11] or 'EUR').strip().upper()
    if len(currency)!=3 or not currency.isalpha(): raise ValueError('La devise doit être un code de trois lettres.')
    available_from=import_date(values[13])
    bio=str(values[14] or '').strip()
    if len(bio)>2000: raise ValueError('La présentation dépasse 2 000 caractères.')
    domains=[]
    for name in import_values(values[5]):
        domain,_=ExpertiseDomain.objects.get_or_create(name=name); domains.append(domain)
    skills=[]
    for name in import_values(values[6]):
        if not domains: raise ValueError('Ajoutez au moins un domaine avant d’importer des compétences.')
        skill,_=Skill.objects.get_or_create(name=name,defaults={'domain':domains[0]}); skills.append(skill)
    certifications=[Certification.objects.get_or_create(name=name)[0] for name in import_values(values[7])]
    verticals=[BusinessVertical.objects.get_or_create(name=name)[0] for name in import_values(values[8])]
    languages=[]
    for item in import_values(values[9]):
        if ':' not in item: raise ValueError('Chaque langue doit avoir le format « Langue:Niveau ».')
        name, level=(part.strip() for part in item.rsplit(':',1)); level=level.upper()
        if not name or level not in dict(ExpertLanguageAssignment.LEVELS): raise ValueError('Une langue ou son niveau est invalide.')
        languages.append((Language.objects.get_or_create(name=name)[0],level))
    expert=Expert.objects.create(first_name=first_name,last_name=last_name,employee_id=employee_id,job_title=job_title,subsidiary=subsidiary,daily_rate=daily_rate,currency=currency,availability=availability,validation_status=validation_status,available_from=available_from,bio=bio)
    expert.domains.set(domains); expert.skills.set(skills); expert.certifications.set(certifications); expert.verticals.set(verticals)
    ExpertLanguageAssignment.objects.bulk_create([ExpertLanguageAssignment(expert=expert,language=language,level=level) for language,level in languages])
    audit(user,'EXPERT_IMPORTED',expert,'Expert importé depuis Excel',metadata={'line':line})
    seen_ids.add(employee_id)
    return expert

@login_required
def expert_import_template(request):
    guard(request.user,['ADMIN_FILIALE','ADMIN_OMEA'])
    if not IMPORT_TEMPLATE.exists(): raise PermissionDenied
    return FileResponse(IMPORT_TEMPLATE.open('rb'),as_attachment=True,filename='modele_import_experts_orbit.xlsx')

@login_required
def expert_import(request):
    guard(request.user,['ADMIN_FILIALE','ADMIN_OMEA'])
    context={'title':'Importer des experts','section':'expert_list','max_rows':500}
    if request.method!='POST': return render(request,'expert_import.html',context)
    upload=request.FILES.get('file')
    if not upload:
        context['upload_error']='Sélectionnez un fichier Excel .xlsx.'; return render(request,'expert_import.html',context)
    if not upload.name.lower().endswith('.xlsx') or upload.size>2*1024*1024:
        context['upload_error']='Le fichier doit être un .xlsx de 2 Mo maximum.'; return render(request,'expert_import.html',context)
    try:
        workbook=load_workbook(upload,read_only=True,data_only=True)
        if 'Experts' not in workbook.sheetnames: raise ValueError('La feuille « Experts » est introuvable.')
        sheet=workbook['Experts']; headers=[str(cell or '').strip() for cell in next(sheet.iter_rows(min_row=4,max_row=4,values_only=True))]
        if headers[:len(IMPORT_HEADERS)] != IMPORT_HEADERS: raise ValueError('Les en-têtes ne correspondent pas au modèle ORBIT.')
    except (InvalidFileException,BadZipFile,OSError,ValueError,StopIteration) as error:
        context['upload_error']=str(error) or 'Le fichier Excel est illisible.'; return render(request,'expert_import.html',context)
    rows=list(sheet.iter_rows(min_row=5,values_only=True))
    if len(rows)>context['max_rows']:
        context['upload_error']=f'Le fichier dépasse la limite de {context["max_rows"]} lignes.'; return render(request,'expert_import.html',context)
    created=[]; errors=[]; seen_ids=set()
    for line,row in enumerate(rows,start=5):
        values=list(row[:len(IMPORT_HEADERS)])
        if not any(value not in (None,'') for value in values): continue
        try:
            with transaction.atomic(): created.append(import_expert_row(values,request.user,line,seen_ids))
        except ValueError as error: errors.append({'line':line,'message':str(error)})
    context.update(created=created,errors=errors,processed=len(created)+len(errors))
    if created: messages.success(request,f'{len(created)} expert(s) importé(s).')
    return render(request,'expert_import.html',context)

@login_required
def request_list(request):
    qs=requests_for(request.user); q=request.GET.get('q','')
    if q: qs=qs.filter(Q(reference__icontains=q)|Q(title__icontains=q))
    if request.GET.get('status') in dict(ExpertRequest.STATUSES): qs=qs.filter(status=request.GET['status'])
    for name in ['subsidiary','domain']:
        if request.GET.get(name,'').isdigit(): qs=qs.filter(**{name+'_id':request.GET[name]})
    return listing(request,qs.order_by('-created_at'),'request_list.html','partials/requests_table.html',title='Demandes d’expertise',section='request_list',statuses=ExpertRequest.STATUSES,**filter_context())

@login_required
def request_edit(request,pk=None):
    guard(request.user,['ADMIN_FILIALE','ADMIN_OMEA'])
    obj=get_object_or_404(requests_for(request.user),pk=pk) if pk else None
    if obj and (not request_writable(request.user,obj) or obj.status!='DRAFT'): raise PermissionDenied
    form=ExpertRequestForm(request.POST or None,instance=obj,user=request.user)
    if request.method=='POST' and form.is_valid():
        with transaction.atomic():
            obj=form.save(commit=False); obj.mode=obj.mode or 'BROADCAST'; obj.requester=obj.requester if pk else request.user; obj.save(); form.save_m2m()
            audit(request.user,'REQUEST_UPDATED' if pk else 'REQUEST_CREATED',obj,'Demande enregistrée en brouillon')
        messages.success(request,'Brouillon enregistré. Vous pouvez maintenant soumettre votre demande.')
        return redirect('request_detail',obj.pk)
    return render(request,'form.html',{'form':form,'title':'Modifier la demande' if pk else 'Nouvelle demande','section':'request_list','back_url':'request_list'})

@login_required
def request_detail(request,pk):
    obj=get_object_or_404(requests_for(request.user),pk=pk)
    if request.method=='POST':
        if not request_writable(request.user,obj): raise PermissionDenied
        try: obj=submit_request(obj,request.user); messages.success(request,'Demande soumise. Le matching peut commencer.')
        except ValueError as e: messages.error(request,str(e))
        return redirect('request_detail',pk)
    labels=['Demande','Matching','Pré-accord','PO','Mission','PV','Évaluation','Clôture']
    current={'DRAFT':0,'SUBMITTED':1,'WAITING_MATCHING':1,'WAITING_PRE_AGREEMENT':2,'WAITING_PO':3,'IN_PROGRESS':4,'WAITING_EVALUATION':6,'CLOSED':7}.get(obj.status,0)
    can_respond_broadcast=(role(request.user)=='ADMIN_FILIALE' and obj.mode=='BROADCAST' and obj.status=='WAITING_MATCHING' and obj.subsidiary_id!=subsidiary_id(request.user))
    responder_experts=Expert.objects.filter(subsidiary_id=subsidiary_id(request.user),is_active=True,validation_status='APPROVED',availability='AVAILABLE').order_by('last_name','first_name') if can_respond_broadcast else Expert.objects.none()
    return render(request,'request_detail.html',{'item':obj,'title':obj.reference,'section':'request_list','steps':[(label,'done' if i<current else 'current' if i==current else '') for i,label in enumerate(labels)],
      'editable':request_writable(request.user,obj) and obj.status=='DRAFT','can_add_po':request_writable(request.user,obj) and obj.status=='WAITING_PO','events':audit_for(request.user).filter(entity_type='ExpertRequest',entity_id=str(pk)),
      'solicitations':solicitations_for(request.user).filter(request=obj),'proposals':obj.proposals.select_related('expert__subsidiary','proposed_by'),'can_respond_broadcast':can_respond_broadcast,'responder_experts':responder_experts,'documents':obj.documents.all()})

@login_required
@require_POST
def purchase_order_upload(request,pk):
    obj=get_object_or_404(requests_for(request.user),pk=pk)
    if not request_writable(request.user,obj): raise PermissionDenied
    try:
        mission=register_purchase_order(obj,request.FILES.get('po_file'),request.user)
        messages.success(request,'Bon de commande ajouté. La mission est désormais planifiée.')
        return redirect('mission_detail',mission.pk)
    except ValueError as error:
        messages.error(request,str(error))
        return redirect('request_detail',pk)

@login_required
@require_POST
def broadcast_proposal_create(request,pk):
    guard(request.user,['ADMIN_FILIALE'])
    obj=get_object_or_404(requests_for(request.user),pk=pk)
    expert=get_object_or_404(Expert.objects.filter(subsidiary_id=subsidiary_id(request.user),is_active=True,validation_status='APPROVED',availability='AVAILABLE'),pk=request.POST.get('expert_id'))
    try:
        create_broadcast_proposal(obj,expert,request.user)
        messages.success(request,f'{expert.full_name} a été proposé à la filiale demandeuse.')
    except ValueError as error:
        messages.error(request,str(error))
    return redirect('request_detail',pk)

@login_required
def solicitation_list(request):
    qs=solicitations_for(request.user)
    q=request.GET.get('q','')
    if q: qs=qs.filter(Q(request__title__icontains=q)|Q(expert__last_name__icontains=q)|Q(request__reference__icontains=q)|Q(subject__icontains=q))
    if request.GET.get('status') in dict(ExpertSolicitation.STATUSES): qs=qs.filter(status=request.GET['status'])
    page=Paginator(qs.order_by('-created_at'),12).get_page(request.GET.get('page'))
    page.object_list=[with_deadline(s) for s in page.object_list]
    params=request.GET.copy(); params.pop('page',None)
    return render(request,'partials/solicitations_table.html' if request.headers.get('HX-Request')=='true' else 'solicitation_list.html',{'page':page,'query':params.urlencode(),'title':'Sollicitations','section':'solicitation_list','statuses':ExpertSolicitation.STATUSES})

@login_required
def solicitation_detail(request,pk):
    s=get_object_or_404(solicitations_for(request.user),pk=pk)
    can_respond=role(request.user)=='ADMIN_FILIALE' and s.expert.subsidiary_id==subsidiary_id(request.user) and s.status=='PENDING' and s.deadline>timezone.now()
    if request.method=='POST':
        if not can_respond: raise PermissionDenied
        try:
            if request.POST.get('action')=='accept': accept_solicitation(s,request.user)
            else: refuse_solicitation(s,request.user,request.POST.get('reason',''))
            messages.success(request,'Votre réponse a été enregistrée.')
        except ValueError as e: messages.error(request,str(e))
        return redirect('solicitation_detail',pk)
    return render(request,'solicitation_detail.html',{'item':with_deadline(s),'can_respond':can_respond,'title':'Pré-accord fournisseur','section':'solicitation_list','events':audit_for(request.user).filter(entity_type='ExpertSolicitation',entity_id=str(pk))})

@login_required
def solicitation_create(request,expert_id):
    guard(request.user,['ADMIN_FILIALE','ADMIN_OMEA'])
    # Une demande ciblée déjà ouverte doit pouvoir être soumise même si la
    # disponibilité de l'expert a changé entre l'affichage du formulaire et le POST.
    # Le service applique ensuite la validation métier et renvoie un message utile.
    expert=get_object_or_404(Expert.objects.select_related('subsidiary').filter(is_active=True,validation_status='APPROVED'),pk=expert_id)
    eligible=requests_for(request.user).filter(status='WAITING_MATCHING')
    if request.method=='POST':
        try:
            with transaction.atomic():
                expert=Expert.objects.select_for_update().get(pk=expert.pk)
                if request.POST.get('request_id'):
                    obj=get_object_or_404(eligible,pk=request.POST.get('request_id'))
                    s=create_solicitation(obj,expert,request.user)
                    messages.success(request,'Sollicitation créée. Le fournisseur dispose de 48 heures.')
                else:
                    s=create_direct_solicitation(expert,request.user,request.POST)
                    messages.success(request,'Sollicitation directe envoyée. La filiale de l’expert dispose de 48 heures pour répondre.')
        except ValueError as error:
            messages.error(request,str(error))
            return redirect('solicitation_create',expert_id)
        response=redirect('solicitation_detail',s.pk)
        if request.headers.get('HX-Request')=='true':
            response=HttpResponse(); response['HX-Redirect']=redirect('solicitation_detail',s.pk).url
        return response
    return render(request,'solicitation_create.html',{'expert':expert,'eligible':eligible,'title':'Créer une demande ciblée','section':'matching'})

@login_required
def mission_list(request):
    qs=missions_for(request.user)
    if request.GET.get('q'): qs=qs.filter(request__title__icontains=request.GET['q'])
    return listing(request,qs.order_by('-start_date'),'mission_list.html','partials/missions_table.html',title='Missions',section='mission_list')

@login_required
def mission_detail(request,pk):
    obj=get_object_or_404(missions_for(request.user),pk=pk)
    can_add_pv=role(request.user)=='ADMIN_FILIALE' and obj.supplier_subsidiary_id==subsidiary_id(request.user) and obj.request.status=='IN_PROGRESS' and not obj.pv_document_id
    can_evaluate=request.user==obj.request.requester and obj.request.status=='WAITING_EVALUATION' and not hasattr(obj,'evaluation')
    return render(request,'mission_detail.html',{'item':obj,'title':obj.request.title,'section':'mission_list','can_add_pv':can_add_pv,'can_evaluate':can_evaluate,'evaluation_form':EvaluationForm(),'events':audit_for(request.user).filter(entity_type='Mission',entity_id=str(pk))})

@login_required
@require_POST
def delivery_report_upload(request,pk):
    mission=get_object_or_404(missions_for(request.user),pk=pk)
    if role(request.user)!='ADMIN_FILIALE' or mission.supplier_subsidiary_id!=subsidiary_id(request.user): raise PermissionDenied
    try:
        register_delivery_report(mission,request.FILES.get('pv_file'),request.user)
        messages.success(request,'PV ajouté. La filiale demandeuse peut maintenant évaluer la mission.')
    except ValueError as error:
        messages.error(request,str(error))
    return redirect('mission_detail',pk)

@login_required
@require_POST
def mission_evaluate(request,pk):
    mission=get_object_or_404(missions_for(request.user),pk=pk)
    if request.user!=mission.request.requester: raise PermissionDenied
    form=EvaluationForm(request.POST)
    if form.is_valid():
        try:
            evaluation=create_evaluation(mission,request.user,form.cleaned_data['rating'],form.cleaned_data['comment'])
            messages.success(request,'Évaluation enregistrée. La mission est clôturée.')
            return redirect('mission_detail',evaluation.mission_id)
        except ValueError as error:
            messages.error(request,str(error))
    else:
        messages.error(request,'Complétez la note et le commentaire avant de valider.')
    return redirect('mission_detail',pk)

@login_required
def notifications(request):
    qs=request.user.notifications.order_by('-created_at')
    if request.GET.get('unread')=='1': qs=qs.filter(is_read=False)
    return listing(request,qs,'notifications.html','partials/notifications_list.html',title='Notifications',section='notifications')

@login_required
def notification_preview(request):
    return render(request,'partials/notifications_dropdown.html',{'items':request.user.notifications.order_by('-created_at')[:5]})

@login_required
@require_POST
def mark_notification_read(request,pk):
    n=get_object_or_404(Notification,pk=pk,user=request.user)
    n.is_read=True; n.save(update_fields=['is_read'])
    if request.headers.get('HX-Request')=='true':
        response=notifications(request); response['HX-Trigger']='notificationsChanged'; return response
    return redirect('notifications')

@login_required
def notification_open(request,pk):
    """Open a notification only while its target remains in the user's scope."""
    notification=get_object_or_404(Notification,pk=pk,user=request.user)
    notification.is_read=True; notification.save(update_fields=['is_read'])
    try:
        match=resolve(notification.url)
        target_pk=match.kwargs.get('pk')
        visible={
            'request_detail': requests_for(request.user).filter(pk=target_pk).exists(),
            'solicitation_detail': solicitations_for(request.user).filter(pk=target_pk).exists(),
            'mission_detail': missions_for(request.user).filter(pk=target_pk).exists(),
        }.get(match.url_name,False)
        if visible: return redirect(notification.url)
    except (Resolver404, TypeError, ValueError):
        pass
    messages.info(request,'Cette notification concerne un élément qui n’est plus disponible dans votre périmètre.')
    return redirect('notifications')

@login_required
def audit_list(request):
    qs=audit_for(request.user)
    if request.GET.get('q'): qs=qs.filter(Q(description__icontains=request.GET['q'])|Q(action__icontains=request.GET['q']))
    return listing(request,qs,'audit_list.html','partials/audit_timeline.html',title='Journal d’activité',section='audit_list')

@login_required
def reports(request):
    reqs=requests_for(request.user); ex=experts_for(request.user); sols=solicitations_for(request.user)
    total=sols.filter(status__in=['ACCEPTED','REFUSED']).count()
    rate=round(sols.filter(status='ACCEPTED').count()/total*100) if total else None
    return render(request,'reports.html',{'title':'Reporting & indicateurs','section':'reports','acceptance':rate,
      'average':ex.aggregate(value=Avg('average_score'))['value'],'by_status':[dict(label=dict(ExpertRequest.STATUSES).get(row['status'],row['status']),total=row['total']) for row in reqs.values('status').annotate(total=Count('id'))],
      'by_subsidiary':ex.values('subsidiary__name').annotate(total=Count('id')),'mission_count':missions_for(request.user).count(),'total_experts':ex.count()})

@login_required
def functional_admin(request):
    guard(request.user,['ADMIN_OMEA'])
    return render(request,'functional_admin.html',{'title':'Administration fonctionnelle','section':'functional_admin','subsidiaries':Subsidiary.objects.select_related('cluster'),'profiles':Profile.objects.select_related('user','subsidiary'),'domains':ExpertiseDomain.objects.annotate(total=Count('expert'))})
