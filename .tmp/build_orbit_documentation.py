from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"C:\Users\hamza\Desktop\mvp")
OUT = ROOT / "docs" / "Documentation_ORBIT.docx"

DARK = "17202B"
ORANGE = "FF7900"
GRAY = "657080"
LIGHT = "F4F5F7"
BORDER = "D9D9D9"

def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)

def set_cell_border(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_border(cell)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    header = t.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        if widths:
            header.cells[i].width = Cm(widths[i])
        shade(header.cells[i], DARK)
        set_cell_text(header.cells[i], text, bold=True, color="FFFFFF", size=8.5)
    for idx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, text in enumerate(row):
            if widths:
                cells[i].width = Cm(widths[i])
            if idx % 2 == 1:
                shade(cells[i], LIGHT)
            set_cell_text(cells[i], text, size=8.5)
    for row in t.rows:
        for cell in row.cells:
            cell.margin_top = Cm(0.08)
            cell.margin_bottom = Cm(0.08)
            cell.margin_left = Cm(0.10)
            cell.margin_right = Cm(0.10)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return t

def paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    for r in p.runs:
        r.font.name = "Arial"; r._element.rPr.rFonts.set(qn("w:ascii"), "Arial"); r._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        r.font.size = Pt(10)
    return p

def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text).font.size = Pt(10)
    return p

def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        r.font.name = "Arial"; r._element.rPr.rFonts.set(qn("w:ascii"), "Arial"); r._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p

doc = Document()
for section in doc.sections:
    section.top_margin = Cm(2.2); section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)

styles = doc.styles
styles["Normal"].font.name = "Arial"; styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial"); styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(10)
for style_name, size in [("Title", 28), ("Heading 1", 17), ("Heading 2", 12)]:
    style = styles[style_name]
    style.font.name = "Arial"; style._element.rPr.rFonts.set(qn("w:ascii"), "Arial"); style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = styles["Title"]._element.get_or_add_pPr()
for title_border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(title_border)

# Cover
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(70); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("ORBIT"); r.bold = True; r.font.size = Pt(20); r.font.name = "Arial"; r.font.color.rgb = RGBColor.from_string(ORANGE)
title = doc.add_paragraph(style="Title"); title.paragraph_format.space_before = Pt(12); title.paragraph_format.space_after = Pt(12)
title.add_run("Documentation fonctionnelle et technique")
subtitle = doc.add_paragraph(); subtitle.paragraph_format.space_after = Pt(28)
subtitle.add_run("Application de pilotage de l expertise Orange Middle East and Africa").italic = True
paragraph(doc, "Cette documentation explique l utilisation de l application ORBIT, son organisation fonctionnelle, ses règles de sécurité et les opérations nécessaires pour l installer et l exploiter.")
table(doc, ["Document", "Valeur"], [["Version", "1.0"], ["Public", "Utilisateurs métiers, administrateurs et équipe technique"], ["Application", "ORBIT"], ["Technologie", "Django 4.2, SQLite ou PostgreSQL, HTMX"], ["Statut", "Documentation de référence du dépôt"]], [4.0, 11.5])
doc.add_page_break()

heading(doc, "Sommaire")
for item in ["1. Présentation et périmètre", "2. Accès et rôles", "3. Parcours métier de la demande", "4. Gestion des experts et import Excel", "5. Administration et référentiels", "6. Architecture technique", "7. Installation et configuration", "8. Exploitation et contrôles", "9. Tests et limites actuelles"]:
    bullet(doc, item)

heading(doc, "1 Présentation et périmètre")
paragraph(doc, "ORBIT centralise les besoins d expertise ICT des filiales et les met en relation avec des experts internes. L application permet de créer une demande, rechercher les profils pertinents, solliciter la filiale qui porte l expert et suivre les décisions jusqu à la clôture de la mission.")
paragraph(doc, "Le périmètre applicatif comprend le référentiel des experts, les demandes, les pré accords, les missions, les notifications, le journal d audit et les indicateurs. Les données sont volontairement filtrées selon le rôle et la filiale de l utilisateur connecté.")
heading(doc, "Objectifs", 2)
for text in ["Réduire le temps nécessaire pour identifier une expertise disponible.", "Conserver une trace des décisions et des changements de statut.", "Sécuriser les échanges entre filiales grâce à un périmètre d accès explicite.", "Produire un reporting exploitable pour les équipes OMEA."]:
    bullet(doc, text)

heading(doc, "2 Accès et rôles")
paragraph(doc, "La connexion se fait sur la page Connexion. Les comptes sont administrés par Django et chaque utilisateur possède un profil ORBIT contenant son rôle et, lorsque nécessaire, sa filiale de rattachement.")
table(doc, ["Rôle", "Accès et actions"], [["Demandeur", "Crée et consulte ses propres demandes. Suit les sollicitations associées. Évalue une mission dont il est le demandeur."], ["Administrateur demandeur", "Gère les demandes de sa filiale, accède au matching et crée les sollicitations."], ["Administrateur fournisseur", "Consulte et gère les experts de sa filiale. Importe les experts. Accepte ou refuse les pré accords qui concernent sa filiale."], ["Administrateur OMEA", "Dispose de la vue transverse, consulte le reporting et l administration, peut gérer les experts de toutes les filiales."]], [4.2, 11.3])
paragraph(doc, "Les contrôles d autorisation sont appliqués côté serveur. Une URL connue ne permet donc pas de contourner le périmètre d une filiale ou le rôle requis.", "Les contrôles d autorisation")

heading(doc, "3 Parcours métier de la demande")
paragraph(doc, "Le parcours visuel d une demande présente huit étapes. Elles permettent aux utilisateurs de se repérer sur l avancement, même lorsque certaines étapes dépendent d opérations administratives complémentaires.")
table(doc, ["Étape", "Statut ou action associée", "Responsable principal"], [["1 Demande", "Création en brouillon puis soumission", "Demandeur ou administrateur demandeur"], ["2 Matching", "Recherche d experts approuvés et disponibles", "Administrateur demandeur ou OMEA"], ["3 Pré accord", "Sollicitation envoyée avec une échéance de 48 heures", "Administrateur fournisseur"], ["4 PO attendu", "Pré accord accepté", "Processus administratif"], ["5 Mission", "Mission planifiée ou en cours", "Équipes métier"], ["6 PV", "Procès verbal attendu", "Processus administratif"], ["7 Évaluation", "Évaluation de la mission par le demandeur", "Demandeur"], ["8 Clôture", "Mission et demande terminées", "Système après évaluation"]], [3.0, 8.0, 4.5])
heading(doc, "Règles de pré accord", 2)
for text in ["Seul un expert approuvé et disponible peut être sollicité.", "La sollicitation passe l expert au statut Pré accord et la demande au statut Pré accord.", "Un refus impose un motif et libère l expert pour un nouveau matching.", "Une réponse acceptée fait avancer la demande à PO attendu.", "Une sollicitation expirée est escaladée et l expert redevient disponible."]:
    bullet(doc, text)
paragraph(doc, "L échéance de 48 heures est traitée par la commande process_timeouts. Les rappels sont générés par send_pending_reminders. Ces commandes doivent être planifiées en production.")

heading(doc, "4 Gestion des experts et import Excel")
paragraph(doc, "La page Experts permet de rechercher et filtrer les profils. Les filtres portent notamment sur la filiale, le cluster, le domaine, les compétences, les certifications, les langues, le tarif, la disponibilité et l évaluation.")
heading(doc, "Création manuelle", 2)
paragraph(doc, "Un administrateur fournisseur ou OMEA peut créer un expert avec son matricule, ses informations professionnelles, sa filiale, ses domaines, ses compétences, ses certifications, ses verticales métier, son tarif, sa disponibilité et sa date de disponibilité.")
heading(doc, "Import Excel", 2)
paragraph(doc, "Depuis Experts, le bouton Importer depuis Excel donne accès à un formulaire et à un modèle de fichier. Le modèle contient les colonnes attendues et deux exemples. Le fichier doit être au format XLSX, peser au plus 2 Mo et contenir au plus 500 lignes de données.")
table(doc, ["Colonne", "Règle"], [["Matricule, Prénom, Nom, Fonction", "Obligatoires."], ["Code filiale", "Obligatoire pour OMEA. Pour un fournisseur, il doit correspondre à sa filiale ou peut être laissé vide."], ["Domaines, Compétences, Certifications, Verticales", "Valeurs séparées par une virgule. Les référentiels absents sont créés."], ["Langues", "Format Langue:Niveau. Niveaux autorisés de A1 à C2."], ["TJM et Devise", "TJM positif ou nul. Devise à trois lettres, EUR par défaut."], ["Disponibilité et statut", "Valeurs ORBIT contrôlées. Un fournisseur crée toujours les experts en brouillon."], ["Présentation et date", "Présentation limitée à 2000 caractères. Date au format AAAA-MM-JJ."]], [5.0, 10.5])
paragraph(doc, "L import traite les lignes indépendamment. Une erreur sur une ligne n annule pas les autres lignes. Le rapport affiche les experts créés et les numéros de lignes rejetées. Chaque création est enregistrée dans le journal d audit.")

heading(doc, "5 Administration et référentiels")
paragraph(doc, "La section Administration est accessible aux administrateurs OMEA. Elle présente les filiales, les profils utilisateurs et les domaines d expertise. Elle est actuellement en lecture seule dans l interface ORBIT.")
paragraph(doc, "Pour créer un cluster ou une filiale, utiliser l administration Django avec un compte super administrateur sur /admin/. Créer le cluster en premier, puis la filiale avec son nom, son code unique et son cluster. Cette séparation évite une filiale sans rattachement organisationnel.")
heading(doc, "Référentiels", 2)
for text in ["Clusters et filiales définissent le périmètre organisationnel.", "Domaines et compétences alimentent le matching et les filtres.", "Certifications, verticales et langues enrichissent les profils experts.", "Les rôles sont initialisés avec la commande seed_roles."]:
    bullet(doc, text)

heading(doc, "6 Architecture technique")
table(doc, ["Composant", "Responsabilité"], [["Django 4.2", "Framework web, authentification, ORM, vues, formulaires et administration."], ["orbit_app models", "Modèles métier : experts, demandes, sollicitations, missions, évaluations, notifications et audit."], ["orbit_app services", "Transitions métier transactionnelles : soumission, sollicitation, réponse, timeout et évaluation."], ["orbit_app selectors", "Filtrage des données selon les rôles et filiales."], ["Templates et HTMX", "Interface responsive, filtres asynchrones et dialogue de sollicitation."], ["SQLite ou PostgreSQL", "Stockage local par défaut ou base de données de production via DATABASE_URL."], ["openpyxl", "Lecture et validation des fichiers Excel importés."]], [4.6, 10.9])
paragraph(doc, "Les assets front end sont servis localement. Les icônes Tabler et HTMX ne dépendent pas d un CDN, ce qui facilite l exécution dans un environnement interne.")

heading(doc, "7 Installation et configuration")
heading(doc, "Installation locale", 2)
for text in ["Créer un environnement virtuel Python.", "Installer les dépendances avec pip install -r requirements.txt.", "Exécuter python manage.py migrate.", "Initialiser les rôles avec python manage.py seed_roles.", "Créer les données de démonstration avec seed_demo_data et seed_design_demo si nécessaire.", "Démarrer le serveur avec python manage.py runserver."]:
    bullet(doc, text)
table(doc, ["Variable", "Usage"], [["DJANGO_SECRET_KEY", "Clé secrète Django. Remplacer impérativement la valeur de développement."], ["DJANGO_DEBUG", "Mettre à 0 en production."], ["DJANGO_ALLOWED_HOSTS", "Liste des hôtes HTTP autorisés."], ["DATABASE_URL", "Active PostgreSQL quand l URL commence par postgres."], ["TIME_ZONE", "Fuseau horaire. Valeur par défaut Africa Casablanca."], ["EMAIL_BACKEND", "Backend e mail. Console par défaut en développement."], ["DEFAULT_FROM_EMAIL", "Adresse expéditrice des notifications e mail."]], [4.8, 10.7])
paragraph(doc, "Les comptes de démonstration sont demo.demandeur, demo.fournisseur et demo.omea. Le mot de passe de démonstration est ORBIT-demo-2026 et ne doit pas être utilisé en production.")

heading(doc, "8 Exploitation et contrôles")
table(doc, ["Commande", "Objectif", "Fréquence recommandée"], [["python manage.py check", "Vérifier la configuration Django.", "Avant chaque déploiement."], ["python manage.py test", "Exécuter les tests automatisés.", "Avant intégration et déploiement."], ["python manage.py process_timeouts", "Escalader les pré accords à échéance.", "Planifiée régulièrement."], ["python manage.py send_pending_reminders", "Créer les rappels sur les pré accords en attente.", "Planifiée régulièrement."], ["python manage.py runserver", "Démarrer le serveur de développement.", "Développement uniquement."]], [5.0, 7.0, 3.5])
paragraph(doc, "Le journal d activité conserve les actions sensibles, notamment les demandes soumises, les sollicitations créées, les réponses fournisseurs, les escalades, les évaluations et les imports Excel. Les notifications sont personnelles et doivent être marquées comme lues par une requête protégée.")

heading(doc, "9 Tests et limites actuelles")
paragraph(doc, "La suite automatisée actuelle couvre les règles de périmètre, les rôles, le rendu des pages, les demandes, le matching, les sollicitations, les refus, le timeout 48 heures, les évaluations et l import Excel. La commande de référence est python manage.py test.")
table(doc, ["Sujet", "État actuel"], [["Demandes et matching", "Opérationnels et testés."], ["Pré accords et timeout", "Opérationnels et testés."], ["Évaluation et clôture", "Opérationnelles et testées."], ["Import Excel experts", "Opérationnel et testé."], ["Administration ORBIT", "Consultation seulement. Création de filiale via Django Admin."], ["Dépôt de PO et PV", "Modèles présents, écran et workflow de dépôt à compléter."], ["Production", "Nécessite des secrets, une base PostgreSQL, un serveur WSGI et la planification des commandes." ]], [5.0, 10.5])
paragraph(doc, "Avant une mise en production, remplacer les secrets de développement, restreindre les hôtes autorisés, configurer le stockage des médias, planifier les tâches de pré accord et valider les politiques d accès avec les administrateurs Orange MEA.")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
